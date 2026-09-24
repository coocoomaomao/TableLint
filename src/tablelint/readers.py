from __future__ import annotations

import csv
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from .models import Finding, ReadResult, Severity, TableData


_TABLE_RE = re.compile(
    r"\\begin\{table\*?\}(.*?)\\end\{table\*?\}",
    re.DOTALL,
)
_TABULAR_RE = re.compile(
    r"\\begin\{tabular\}\s*\{([^{}]*)\}(.*?)\\end\{tabular\}",
    re.DOTALL,
)
_CAPTION_RE = re.compile(r"\\caption(?:\s*\[[^\]]*\])?\s*\{")
_LABEL_RE = re.compile(r"\\label\s*\{")
_ROW_BREAK_RE = re.compile(r"(?<!\\)\\\\(?![A-Za-z])")
_RULE_RE = re.compile(
    r"\\(?:hline|toprule|midrule|bottomrule|addlinespace)"
    r"|\\cline\s*\{[^{}]*\}"
)
_UNESCAPED_AMP_RE = re.compile(r"(?<!\\)&")


def _line_number(text: str, offset: int) -> int:
    return text.count("\n", 0, offset) + 1


def _strip_latex_comments(text: str) -> str:
    out: list[str] = []
    for line in text.splitlines(keepends=True):
        cut = None
        for idx, char in enumerate(line):
            if char != "%":
                continue
            backslashes = 0
            pos = idx - 1
            while pos >= 0 and line[pos] == "\\":
                backslashes += 1
                pos -= 1
            if backslashes % 2 == 0:
                cut = idx
                break
        if cut is None:
            out.append(line)
        else:
            ending = "\n" if line.endswith("\n") else ""
            out.append(line[:cut] + ending)
    return "".join(out)


def _trim_matrix(matrix: list[list[Any]]) -> list[list[Any]]:
    def blank(value: Any) -> bool:
        return value is None or (isinstance(value, str) and not value.strip())

    while matrix and all(blank(value) for value in matrix[-1]):
        matrix.pop()

    if not matrix:
        return []

    width = max(len(row) for row in matrix)
    while width > 0 and all(
        blank(row[width - 1] if width - 1 < len(row) else None)
        for row in matrix
    ):
        width -= 1

    return [row[:width] for row in matrix]


def read_csv(path: Path) -> ReadResult:
    result = ReadResult()
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            matrix = [list(row) for row in csv.reader(handle, dialect)]
    except (OSError, UnicodeError, csv.Error) as exc:
        result.findings.append(
            Finding(
                path=path,
                severity=Severity.ERROR,
                code="FILE_UNREADABLE",
                message=f"Could not read CSV file: {exc}.",
            )
        )
        return result

    matrix = _trim_matrix(matrix)
    if not matrix:
        result.findings.append(
            Finding(path, Severity.WARNING, "EMPTY_TABLE", "CSV file contains no table data.")
        )
        return result

    result.tables.append(
        TableData(
            path=path,
            name=path.stem,
            headers=matrix[0],
            rows=matrix[1:],
            source_kind="csv",
        )
    )
    return result


def read_xlsx(path: Path) -> ReadResult:
    result = ReadResult()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        result.findings.append(
            Finding(
                path=path,
                severity=Severity.ERROR,
                code="FILE_UNREADABLE",
                message=f"Could not read XLSX workbook: {exc}.",
            )
        )
        return result

    try:
        for sheet in workbook.worksheets:
            matrix = _trim_matrix([list(row) for row in sheet.iter_rows(values_only=True)])
            if not matrix:
                result.findings.append(
                    Finding(
                        path=path,
                        severity=Severity.INFO,
                        code="EMPTY_SHEET",
                        message=f"Worksheet '{sheet.title}' contains no table data.",
                        table=sheet.title,
                    )
                )
                continue
            result.tables.append(
                TableData(
                    path=path,
                    name=sheet.title,
                    headers=matrix[0],
                    rows=matrix[1:],
                    source_kind="xlsx",
                )
            )
    finally:
        workbook.close()

    return result


def _parse_latex_rows(body: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw in _ROW_BREAK_RE.split(body):
        cleaned = _RULE_RE.sub("", raw).strip()
        if not cleaned:
            continue
        if "\\multicolumn" in cleaned or "\\multirow" in cleaned:
            # Width inference is ambiguous; keep the row for basic checks but skip
            # it in the structural-width comparison.
            cells = [cell.strip() for cell in _UNESCAPED_AMP_RE.split(cleaned)]
            rows.append(cells)
            continue
        rows.append([cell.strip() for cell in _UNESCAPED_AMP_RE.split(cleaned)])
    return rows


def read_latex(path: Path) -> ReadResult:
    result = ReadResult()
    try:
        text = path.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        result.findings.append(
            Finding(
                path=path,
                severity=Severity.ERROR,
                code="FILE_UNREADABLE",
                message=f"Could not read LaTeX file: {exc}.",
            )
        )
        return result

    clean = _strip_latex_comments(text)
    covered_spans: list[tuple[int, int]] = []
    table_index = 0

    for table_match in _TABLE_RE.finditer(clean):
        covered_spans.append(table_match.span())
        table_index += 1
        table_body = table_match.group(1)
        table_name = f"table-{table_index}"
        line = _line_number(clean, table_match.start())

        if not _CAPTION_RE.search(table_body):
            result.findings.append(
                Finding(
                    path,
                    Severity.WARNING,
                    "LATEX_CAPTION_MISSING",
                    "LaTeX table environment has no caption.",
                    table=table_name,
                    line=line,
                )
            )
        if not _LABEL_RE.search(table_body):
            result.findings.append(
                Finding(
                    path,
                    Severity.INFO,
                    "LATEX_LABEL_MISSING",
                    "LaTeX table environment has no label.",
                    table=table_name,
                    line=line,
                )
            )

        tabular_match = _TABULAR_RE.search(table_body)
        if tabular_match is None:
            result.findings.append(
                Finding(
                    path,
                    Severity.WARNING,
                    "LATEX_TABULAR_MISSING",
                    "LaTeX table environment contains no tabular environment.",
                    table=table_name,
                    line=line,
                )
            )
            continue

        rows = _parse_latex_rows(tabular_match.group(2))
        if not rows:
            result.findings.append(
                Finding(
                    path,
                    Severity.WARNING,
                    "EMPTY_TABLE",
                    "LaTeX tabular contains no data rows.",
                    table=table_name,
                    line=line,
                )
            )
            continue

        simple_widths = [
            len(row)
            for row in rows
            if not any("\\multicolumn" in cell or "\\multirow" in cell for cell in row)
        ]
        if len(set(simple_widths)) > 1:
            result.findings.append(
                Finding(
                    path,
                    Severity.WARNING,
                    "LATEX_ROW_WIDTH_MISMATCH",
                    f"LaTeX tabular rows have inconsistent cell counts: {sorted(set(simple_widths))}.",
                    table=table_name,
                    line=line,
                )
            )

        result.tables.append(
            TableData(
                path=path,
                name=table_name,
                headers=rows[0],
                rows=rows[1:],
                source_kind="latex",
                metadata={"line": line},
            )
        )

    def inside_table(start: int) -> bool:
        return any(left <= start < right for left, right in covered_spans)

    standalone_index = 0
    for tabular_match in _TABULAR_RE.finditer(clean):
        if inside_table(tabular_match.start()):
            continue
        standalone_index += 1
        name = f"tabular-{standalone_index}"
        line = _line_number(clean, tabular_match.start())
        rows = _parse_latex_rows(tabular_match.group(2))
        if not rows:
            continue
        simple_widths = [
            len(row)
            for row in rows
            if not any("\\multicolumn" in cell or "\\multirow" in cell for cell in row)
        ]
        if len(set(simple_widths)) > 1:
            result.findings.append(
                Finding(
                    path,
                    Severity.WARNING,
                    "LATEX_ROW_WIDTH_MISMATCH",
                    f"Standalone tabular rows have inconsistent cell counts: {sorted(set(simple_widths))}.",
                    table=name,
                    line=line,
                )
            )
        result.tables.append(
            TableData(
                path=path,
                name=name,
                headers=rows[0],
                rows=rows[1:],
                source_kind="latex",
                metadata={"line": line},
            )
        )

    if not result.tables and not result.findings:
        result.findings.append(
            Finding(
                path,
                Severity.INFO,
                "NO_TABLES_FOUND",
                "No LaTeX table/tabular environments were found.",
            )
        )

    return result


def read_file(path: Path) -> ReadResult:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    if suffix == ".tex":
        return read_latex(path)

    return ReadResult(
        findings=[
            Finding(
                path,
                Severity.ERROR,
                "UNSUPPORTED_FILE",
                f"Unsupported table file type: {suffix or '<none>'}.",
            )
        ]
    )
