from pathlib import Path

from openpyxl import Workbook

from tablelint.readers import read_csv, read_latex, read_xlsx


def test_read_csv(tmp_path: Path) -> None:
    path = tmp_path / "results.csv"
    path.write_text("group,mean,p-value\nA,1.2,0.04\n", encoding="utf-8")

    result = read_csv(path)

    assert result.findings == []
    assert len(result.tables) == 1
    assert result.tables[0].headers == ["group", "mean", "p-value"]
    assert result.tables[0].rows[0] == ["A", "1.2", "0.04"]


def test_read_xlsx_multiple_sheets(tmp_path: Path) -> None:
    path = tmp_path / "results.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet.append(["group", "mean"])
    sheet.append(["A", 1.2])
    empty = workbook.create_sheet("Empty")
    workbook.save(path)

    result = read_xlsx(path)

    assert len(result.tables) == 1
    assert result.tables[0].name == "Main"
    assert any(f.code == "EMPTY_SHEET" for f in result.findings)


def test_read_latex_structural_findings(tmp_path: Path) -> None:
    path = tmp_path / "table.tex"
    path.write_text(
        r"""\begin{table}
\begin{tabular}{lll}
A & B & C \\
1 & 2 \\
\end{tabular}
\end{table}
""",
        encoding="utf-8",
    )

    result = read_latex(path)
    codes = {finding.code for finding in result.findings}

    assert "LATEX_CAPTION_MISSING" in codes
    assert "LATEX_LABEL_MISSING" in codes
    assert "LATEX_ROW_WIDTH_MISMATCH" in codes
    assert len(result.tables) == 1


def test_latex_comments_do_not_create_fake_rows(tmp_path: Path) -> None:
    path = tmp_path / "table.tex"
    path.write_text(
        r"""\begin{table}
\caption{Example}
\label{tab:example}
\begin{tabular}{ll}
A & B \\
1 & 2 \\ % comment & fake
\end{tabular}
\end{table}
""",
        encoding="utf-8",
    )

    result = read_latex(path)

    assert "LATEX_ROW_WIDTH_MISMATCH" not in {f.code for f in result.findings}
